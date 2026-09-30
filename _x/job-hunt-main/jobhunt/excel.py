"""Excel workbook: a ranked shortlist sheet plus an unfiltered everything sheet.

The second sheet exists so a scoring mistake is recoverable — you can sort
the full set yourself and see exactly why anything was skipped.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import pipeline

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HOT = PatternFill("solid", fgColor="FEF3C7")     # 75+
GOOD = PatternFill("solid", fgColor="DCFCE7")    # 65-74
DROPPED = PatternFill("solid", fgColor="F3F4F6")

SHORTLIST_COLS = [
    ("score", 7), ("company", 18), ("title", 54), ("location", 30),
    ("bucket", 12), ("H-1B approvals", 15), ("locations", 10),
    ("first seen", 17), ("posted", 12), ("link", 46),
    ("tailor id", 12), ("applied?", 10), ("notes", 30),
]
ALL_COLS = [
    ("score", 7), ("skipped because", 26), ("company", 18), ("title", 54),
    ("location", 30), ("source", 15), ("bucket", 12), ("first seen", 17),
    ("link", 46),
]


def _style_header(sheet, columns: list[tuple[str, int]]) -> None:
    for index, (name, width) in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(columns))}{sheet.max_row or 1}"
    )


def _link(sheet, row: int, column: int, url: str) -> None:
    cell = sheet.cell(row=row, column=column, value=url[:250])
    if url:
        cell.hyperlink = url
        cell.font = Font(color="2563EB", underline="single", size=10)


def render(conn: sqlite3.Connection, path: Path, min_score: float = 55) -> Path:
    book = Workbook()

    # --- Sheet 1: the ranked shortlist -----------------------------------
    sheet = book.active
    sheet.title = "Shortlist"
    rows = pipeline.shortlist(conn, limit=5000, min_score=min_score)
    for index, row in enumerate(rows, start=2):
        sheet.cell(row=index, column=1, value=round(row["score"]))
        sheet.cell(row=index, column=2, value=row["company"])
        sheet.cell(row=index, column=3, value=row["title"])
        sheet.cell(row=index, column=4, value=(row["location"] or "")[:120])
        sheet.cell(row=index, column=5, value=row["bucket"])
        sheet.cell(row=index, column=6, value=row.get("h1b_approvals") or 0)
        sheet.cell(row=index, column=7, value=row.get("duplicate_count", 1))
        sheet.cell(row=index, column=8, value=(row["first_seen"] or "")[:16])
        sheet.cell(row=index, column=9, value=(row["posted_at"] or "")[:10])
        _link(sheet, index, 10, row["url"] or "")
        sheet.cell(row=index, column=11, value=row["uid"][:8])
        fill = HOT if row["score"] >= 75 else GOOD if row["score"] >= 65 else None
        if fill:
            for column in range(1, 4):
                sheet.cell(row=index, column=column).fill = fill
    _style_header(sheet, SHORTLIST_COLS)

    # --- Sheet 2: everything, unfiltered ---------------------------------
    everything = book.create_sheet("All postings")
    all_rows = pipeline.all_rows(conn)
    for index, row in enumerate(all_rows, start=2):
        everything.cell(row=index, column=1, value=round(row["score"] or 0))
        reason = everything.cell(row=index, column=2, value=row["reason"] or "on shortlist")
        everything.cell(row=index, column=3, value=row["company"])
        everything.cell(row=index, column=4, value=row["title"])
        everything.cell(row=index, column=5, value=(row["location"] or "")[:120])
        everything.cell(row=index, column=6, value=row["source"])
        everything.cell(row=index, column=7, value=row["bucket"])
        everything.cell(row=index, column=8, value=(row["first_seen"] or "")[:16])
        _link(everything, index, 9, row["url"] or "")
        if row["reason"]:
            reason.fill = DROPPED
    _style_header(everything, ALL_COLS)

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
