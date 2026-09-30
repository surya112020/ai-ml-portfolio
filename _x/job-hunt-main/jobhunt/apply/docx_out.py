"""Write the tailored résumé as a .docx you can actually upload.

Formatting is deliberately plain. Résumé parsers routinely mangle tables,
text boxes, multiple columns, and anything in a page header — a name placed
in a Word header is a common reason a parsed application comes back with a
blank name field. So: one column, no tables, contact details in the body,
real bullet-list styles, and one standard font.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from .tailor import Tailored

FONT = "Calibri"
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0x99, 0x99, 0x99)


def _base_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing = 1.06

    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)


def _heading(doc: Document, text: str) -> None:
    """Section heading with a bottom rule, drawn as a paragraph border.

    A one-cell table would look the same and break parsers.
    """
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(11)
    para.paragraph_format.space_after = Pt(4)
    run = para.add_run(text.upper())
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = FONT
    para.paragraph_format.tab_stops  # noqa: B018 - touch to ensure pPr exists

    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "999999")
    borders.append(bottom)
    para._p.get_or_add_pPr().append(borders)


def _as_text(value: object) -> str:
    """Tolerate a YAML mapping where a string was meant.

    An unquoted "Name: detail" entry parses as {Name: detail}, which is easy
    to reintroduce when editing resume.yml by hand.
    """
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items())
    return str(value)


def _bullet(doc: Document, text: object) -> None:
    para = doc.add_paragraph(_as_text(text).strip(), style="List Bullet")
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.left_indent = Inches(0.22)
    for run in para.runs:
        run.font.size = Pt(10.5)
        run.font.name = FONT


def render(t: Tailored, resume: dict, profile: dict, path: Path) -> Path:
    ident = profile.get("identity", {})
    doc = Document()
    _base_styles(doc)

    # --- name + contact, in the body so parsers find them ---
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    run = name.add_run(ident.get("name", ""))
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.name = FONT

    contact_bits = [ident.get("base_location"), ident.get("phone"),
                    ident.get("email"), ident.get("linkedin")]
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    crun = contact.add_run("  |  ".join(x for x in contact_bits if x))
    crun.font.size = Pt(9.5)
    crun.font.color.rgb = MUTED
    crun.font.name = FONT

    # --- summary ---
    _heading(doc, "Summary")
    summary = doc.add_paragraph(resume["summary_base"].strip())
    summary.paragraph_format.space_after = Pt(2)
    for r in summary.runs:
        r.font.size = Pt(10.5)
        r.font.name = FONT

    # --- experience, grouped by role, tailored bullets only ---
    _heading(doc, "Experience")
    chosen: dict[str, list[str]] = {}
    for company, _title, text in t.bullets:
        chosen.setdefault(company, []).append(text)
    if not chosen:  # no JD to rank against — fall back to everything
        chosen = {role["company"]: [b["text"] for b in role["bullets"]]
                  for role in resume["experience"]}

    for role in resume["experience"]:
        picked = chosen.get(role["company"])
        if not picked:
            continue
        line = doc.add_paragraph()
        line.paragraph_format.space_before = Pt(7)
        line.paragraph_format.space_after = Pt(1)
        title_run = line.add_run(f"{role['title']} — {role['company']}")
        title_run.font.bold = True
        title_run.font.size = Pt(11)
        title_run.font.name = FONT

        meta = doc.add_paragraph()
        meta.paragraph_format.space_after = Pt(3)
        span = f"{role['start']} – {role.get('end') or 'Present'}"
        mrun = meta.add_run(f"{role['location']}  ·  {span}")
        mrun.font.size = Pt(9.5)
        mrun.font.italic = True
        mrun.font.color.rgb = MUTED
        mrun.font.name = FONT

        for text in picked:
            _bullet(doc, text)

    # --- skills, groups the posting mentions first ---
    _heading(doc, "Technical Skills")
    ordered = t.skill_order + [g for g in resume["skills"]
                               if g not in t.skill_order]
    for group in ordered:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        label = para.add_run(f"{group}: ")
        label.font.bold = True
        label.font.size = Pt(10)
        label.font.name = FONT
        body = para.add_run(", ".join(resume["skills"][group]))
        body.font.size = Pt(10)
        body.font.name = FONT

    # --- education / certifications ---
    _heading(doc, "Education")
    for school in resume["education"]:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        head = para.add_run(f"{school['school']} — ")
        head.font.bold = True
        head.font.size = Pt(10.5)
        head.font.name = FONT
        rest = para.add_run(
            f"{school['degree']}, {school['location']} ({school['ended']})")
        rest.font.size = Pt(10.5)
        rest.font.name = FONT

    _heading(doc, "Certifications")
    for cert in resume["certifications"]:
        _bullet(doc, cert)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
